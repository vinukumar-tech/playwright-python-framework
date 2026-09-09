from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


REPORT_DIR = Path("reports")
EXCEL_FILE = REPORT_DIR / "test_execution.xlsx"


def save_test_result(data, sheet_name="Test Execution"):

    REPORT_DIR.mkdir(exist_ok=True)

    if not EXCEL_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name

        headers = list(data.keys())

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        workbook.save(EXCEL_FILE)

    workbook = load_workbook(EXCEL_FILE)

    if sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.create_sheet(sheet_name)

        headers = list(data.keys())

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    row = sheet.max_row + 1

    headers = [cell.value for cell in sheet[1]]

    for column, header in enumerate(headers, start=1):
        value = data.get(header, "")
        cell = sheet.cell(row=row, column=column, value=value)
        cell.alignment = Alignment(vertical="center")

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, header in enumerate(headers, start=1):

        if header in ["Overall Result", "Request Status"]:

            value = sheet.cell(row=row, column=index).value

            if value in ["PASSED", "In-Approval"]:
                sheet.cell(row=row, column=index).fill = PatternFill(
                    fill_type="solid",
                    fgColor="C6EFCE"
                )

            elif value == "FAILED":
                sheet.cell(row=row, column=index).fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFC7CE"
                )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row_cells in sheet.iter_rows():
        for cell in row_cells:
            cell.border = thin_border

    for column_cells in sheet.columns:

        max_length = 0

        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:

            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        sheet.column_dimensions[column_letter].width = min(max_length + 3, 40)

    sheet.freeze_panes = "A2"

    workbook.save(EXCEL_FILE)

    print(f"Excel report updated: {EXCEL_FILE}")
    print(f"Sheet updated: {sheet_name}")